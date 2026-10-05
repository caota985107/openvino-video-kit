#!/usr/bin/env python3
"""01_extract_ppt.py — PPTX → slides.json

抽出每一頁的標題、內文（含表格）、講者備忘稿，輸出 JSON：
[{"slide":1, "title":"...", "bullets":["..."], "notes":"..."}, ...]

用法:
    python scripts/01_extract_ppt.py deck.pptx [slides.json]
"""
import json
import sys
from pathlib import Path

from pptx import Presentation


def shape_lines(shape):
    """回傳 shape 內所有可讀文字行（含表格）。"""
    lines = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if text:
                lines.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return lines


def extract(src: Path):
    prs = Presentation(str(src))
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None

        bullets = []
        for shape in slide.shapes:
            lines = shape_lines(shape)
            if not lines:
                continue
            if title_shape is not None and shape == title_shape:
                title = " ".join(lines)
            else:
                bullets.extend(lines)

        notes = ""
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame is not None:
                notes = notes_frame.text.strip()

        slides.append(
            {"slide": idx, "title": title, "bullets": bullets, "notes": notes}
        )
    return slides


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "deck.pptx")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "slides.json")
    if not src.exists():
        sys.exit(f"找不到檔案: {src}")
    slides = extract(src)
    out.write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK: {len(slides)} 頁 → {out}")


if __name__ == "__main__":
    main()
